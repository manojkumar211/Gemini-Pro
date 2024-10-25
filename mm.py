import streamlit as st
import os
import pandas as pd
import openpyxl as op
import google.generativeai as genai
from dotenv import load_dotenv
import json
from pypdf import PdfReader
from datetime import datetime


# Load environment variables
load_dotenv()

# Configure the Google Generative AI API key
genai.configure(api_key="AIzaSyDQfOqzMq5syEPdjvuAbP2ffN7V0SNBIOA")

# Load the Gemini Pro model
model = genai.GenerativeModel("gemini-pro")

# Set the page configuration
st.set_page_config(page_title="Requirements")
st.header("Gemini LLM Model")

# User input and file uploader
input_text = st.text_input("Input:", key="input")
upload_file = st.file_uploader("Upload a file", type=['xlsx', 'docx', 'pdf'])
submit = st.button("Submit Here")

def get_gen_response(input_text, file_content):
    # Generate response based on input and file content
    if input_text:
        response = model.generate_content([input_text, file_content])
        return response.txt
    return None

def serialize_value(value):
    if isinstance(value, datetime):
        return value.isoformat()  # Convert datetime to ISO format
    elif isinstance(value, (str, int, float)):
        return value  # Return other basic types as is
    return str(value)  # Convert anything else to string

def process_uploaded_file(upload_file):
    if upload_file.name.endswith('.xlsx'):
        df = op.load_workbook(upload_file)  # Load the workbook
        sheet = df.active
        data = []
        headers = [cell.value for cell in sheet[1]]  # Get headers from the first row

        # Iterate through the rows in the sheet (starting from the second row)
        for row in sheet.iter_rows(min_row=2, values_only=True):
            row_data = {headers[i]: serialize_value(row[i]) for i in range(len(headers))}
            data.append(row_data)

        # Write the data to a JSON file
        json_file_path = "C:/Projects/astria/xlsx_json.json"
        with open(json_file_path, 'w') as json_file:
            json.dump(data, json_file, indent=4)
        return json.dumps(data)  # Return JSON string
    
    elif upload_file.name.endswith('.docx'):
        from docx import Document
        doc = Document(upload_file)
        data = []
        
        # Extract text from each paragraph in the .docx file
        for para in doc.paragraphs:
            if para.text:  # Check if paragraph text is not empty
                data.append({"text": para.text})  # Store each paragraph as a dictionary

        # Write the data to a JSON file
        json_file_path = "C:/Projects/astria/docx_json.json"  # Specify the correct JSON file path
        with open(json_file_path, 'w', encoding='utf-8') as json_file:  # Use encoding='utf-8' to handle special characters
            json.dump(data, json_file, indent=4)
        
        return json.dumps(data)
    
    elif upload_file.name.endswith('.pdf'):
        pdf_reader = PdfReader(upload_file)
        data = []
    
        for page in pdf_reader.pages:
            text = page.extract_text()
            if text:  # Check if text extraction was successful
                data.append({"text": text})  # Store each page's text as a dictionary

        # Write the data to a JSON file
        json_filepdf_path = "C:/Projects/astria/pdf_json.json"  # Specify the correct JSON file path
        with open(json_filepdf_path, 'w', encoding='utf-8') as json_file:  # Use encoding='utf-8' to handle special characters
            json.dump(data, json_file, indent=4)
        
        return json.dumps(data)

    else:
        st.error("Unsupported file format!")
        return None

# Handle submission
if submit:
    file_content = process_uploaded_file(upload_file)
    if file_content:
        response = get_gen_response(input_text, file_content)
        if response:
            st.subheader("The Final Answer")
            st.write(response)